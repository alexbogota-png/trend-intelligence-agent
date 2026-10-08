import base64
import hashlib
import html
import io
import json
import os
from pathlib import Path


def _clean_text(value):
    return "\n".join(line.strip() for line in str(value or "").splitlines() if line.strip())


def extract_presentation(data: bytes, filename: str) -> dict:
    """Extract editable slide content, chart/table data, notes and original images."""
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    assets = {}
    slides = []
    for index, slide in enumerate(presentation.slides, 1):
        item = {"number": index, "title": "", "texts": [], "tables": [], "charts": [], "images": [], "notes": ""}
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = _clean_text(shape.text)
                if text:
                    item["texts"].append(text)
            if getattr(shape, "has_table", False):
                item["tables"].append([
                    [_clean_text(cell.text) for cell in row.cells]
                    for row in shape.table.rows
                ])
            if getattr(shape, "has_chart", False):
                chart = shape.chart
                series = []
                for chart_series in chart.series:
                    series.append({"name": str(chart_series.name or "Serie"), "values": [
                        value if isinstance(value, (str, int, float, bool)) or value is None else str(value)
                        for value in chart_series.values
                    ]})
                item["charts"].append({"type": str(chart.chart_type), "series": series})
            if getattr(shape, "shape_type", None) == 13:
                image = shape.image
                content = image.blob
                digest = hashlib.sha256(content).hexdigest()
                if digest not in assets:
                    assets[digest] = {
                        "mime": image.content_type or "image/png",
                        "data": base64.b64encode(content).decode("ascii"),
                        "width": image.size[0],
                        "height": image.size[1],
                    }
                item["images"].append(digest)
        try:
            item["notes"] = _clean_text(slide.notes_slide.notes_text_frame.text)
        except Exception:
            item["notes"] = ""
        item["title"] = _clean_text(slide.shapes.title.text) if slide.shapes.title else ""
        slides.append(item)

    if not slides:
        raise ValueError("La presentación no contiene diapositivas.")
    return {
        "filename": Path(filename).name,
        "slide_count": len(slides),
        "image_count": sum(len(slide["images"]) for slide in slides),
        "unique_image_count": len(assets),
        "table_count": sum(len(slide["tables"]) for slide in slides),
        "chart_count": sum(len(slide["charts"]) for slide in slides),
        "slides": slides,
        "assets": assets,
    }


def _fallback_analysis(deck: dict) -> dict:
    first_slide = deck["slides"][0]
    title = first_slide.get("title") or next(iter(first_slide.get("texts", [])), "Presentación analizada")
    return {
        "executive_summary": f"Se extrajo el contenido de {deck['slide_count']} diapositivas. El reporte conserva los textos, tablas, gráficos detectables y las imágenes originales para revisión.",
        "kpis": [],
        "findings": [],
        "actions": [],
        "limitations": ["No se generó una interpretación con IA en este proceso. El reporte muestra el contenido extraído y las imágenes de origen."],
        "title": title,
    }


def _summarize(deck: dict) -> dict:
    if not os.getenv("OPENAI_API_KEY"):
        return _fallback_analysis(deck)
    try:
        from openai import OpenAI

        source_slides = []
        total_chars = 0
        for slide in deck["slides"]:
            text = "\n".join(slide["texts"])
            charts = json.dumps(slide["charts"], ensure_ascii=False, default=str)
            tables = json.dumps(slide["tables"], ensure_ascii=False, default=str)
            notes = slide["notes"]
            entry = {"slide": slide["number"], "title": slide["title"], "text": text[:9000], "tables": tables[:3500], "charts": charts[:2500], "notes": notes[:2500]}
            size = sum(len(str(value)) for value in entry.values())
            if total_chars + size > 55_000:
                entry["text"] = entry["text"][:max(0, 55_000 - total_chars)]
                entry["tables"] = ""
                entry["charts"] = ""
                entry["notes"] = ""
            total_chars += sum(len(str(value)) for value in entry.values())
            source_slides.append(entry)
        prompt = (
            "Analiza esta presentación completa y redacta un resumen ejecutivo en español. "
            "Usa exclusivamente el contenido entregado. No inventes datos ni atribuyas causalidad. "
            "Separa hechos de interpretación. Cada hallazgo debe indicar sus diapositivas de evidencia. "
            "No afirmes que inspeccionaste imágenes: usa solo textos, tablas, gráficos y notas extraídos. "
            "Conserva unidades y contexto de los indicadores. Omite campos sin evidencia. "
            "Responde como JSON válido con: title (string), executive_summary (string), "
            "kpis (array de objetos {label, value, context, slide}), "
            "findings (array de objetos {title, detail, evidence, slides}), "
            "actions (array de strings, solo si la evidencia sostiene una recomendación), "
            "limitations (array de strings). Limita a 6 hallazgos y 5 acciones. "
            "Presentación: " + json.dumps(source_slides, ensure_ascii=False, default=str)
        )
        response = OpenAI().chat.completions.create(
            model=os.getenv("OPENAI_MODEL", "gpt-5-mini"),
            messages=[{"role": "user", "content": prompt}],
            response_format={"type": "json_object"},
            max_completion_tokens=2500,
        )
        parsed = json.loads(response.choices[0].message.content or "{}")
        slide_numbers = {slide["number"] for slide in deck["slides"]}
        def valid_list(key, maximum):
            value = parsed.get(key)
            return value[:maximum] if isinstance(value, list) else []
        kpis = []
        for item in valid_list("kpis", 8):
            if not isinstance(item, dict):
                continue
            try:
                source_slide = int(item.get("slide"))
            except (TypeError, ValueError):
                continue
            if source_slide not in slide_numbers:
                continue
            kpis.append({key: str(item.get(key, "")) for key in ("label", "value", "context")} | {"slide": source_slide})
        findings = []
        for item in valid_list("findings", 6):
            if not isinstance(item, dict):
                continue
            evidence_slides = []
            for number in item.get("slides", []):
                try:
                    number = int(number)
                except (TypeError, ValueError):
                    continue
                if number in slide_numbers and number not in evidence_slides:
                    evidence_slides.append(number)
            if evidence_slides:
                findings.append({"title": str(item.get("title", "Hallazgo")), "detail": str(item.get("detail", "")), "evidence": str(item.get("evidence", "")), "slides": evidence_slides})
        return {
            "title": str(parsed.get("title") or deck["slides"][0].get("title") or Path(deck["filename"]).stem),
            "executive_summary": str(parsed.get("executive_summary") or "El archivo se procesó, pero no se pudo generar un resumen ejecutivo con la información disponible."),
            "kpis": kpis,
            "findings": findings,
            "actions": [str(item) for item in valid_list("actions", 5) if str(item).strip()],
            "limitations": [str(item) for item in valid_list("limitations", 5) if str(item).strip()],
        }
    except Exception:
        return _fallback_analysis(deck)


def _slide_html(slide: dict, assets: dict) -> str:
    number = slide["number"]
    content = []
    if slide["texts"]:
        content.append("<div class=\"slide-copy\"><h3>Contenido</h3><ul>" + "".join(f"<li>{html.escape(text)}</li>" for text in slide["texts"]) + "</ul></div>")
    for table in slide["tables"]:
        rows = "".join("<tr>" + "".join(f"<td>{html.escape(str(cell))}</td>" for cell in row) + "</tr>" for row in table)
        content.append(f"<div class=\"slide-data\"><h3>Tabla de la diapositiva</h3><div class=\"table-scroll\"><table>{rows}</table></div></div>")
    for chart in slide["charts"]:
        details = []
        for series in chart["series"]:
            vals = ", ".join(html.escape(str(value)) for value in series["values"])
            details.append(f"<p><strong>{html.escape(series['name'])}</strong>: {vals}</p>")
        content.append(f"<div class=\"slide-data\"><h3>Gráfico · {html.escape(chart['type'])}</h3>{''.join(details)}</div>")
    if slide["notes"]:
        content.append(f"<div class=\"slide-data\"><h3>Notas del presentador</h3><p>{html.escape(slide['notes'])}</p></div>")
    images = []
    for digest in slide["images"]:
        asset = assets[digest]
        images.append(f"<figure><img loading=\"lazy\" src=\"data:{html.escape(asset['mime'])};base64,{asset['data']}\" alt=\"Imagen extraída de la diapositiva {number}\"><figcaption>Imagen original · diapositiva {number} · {asset['width']} × {asset['height']} px</figcaption></figure>")
    if images:
        content.append(f"<div class=\"slide-visuals\"><h3>Imágenes originales</h3><div class=\"image-grid\">{''.join(images)}</div></div>")
    if not content:
        content.append("<p class=\"muted\">No se encontró texto editable ni recurso compatible para esta diapositiva.</p>")
    title = slide["title"] or f"Diapositiva {number}"
    return f"<article class=\"slide\" id=\"slide-{number}\"><div class=\"slide-index\">DIAPOSITIVA {number:02d}</div><h2>{html.escape(title)}</h2>{''.join(content)}</article>"


def build_report_html(deck: dict) -> str:
    analysis = _summarize(deck)
    title = analysis.get("title") or Path(deck["filename"]).stem
    cover_path = Path(__file__).resolve().parent.parent / "static" / "report-cover.jpg"
    cover_data = base64.b64encode(cover_path.read_bytes()).decode("ascii") if cover_path.is_file() else ""
    cover_html = (
        f'<img class="brand-cover" src="data:image/jpeg;base64,{cover_data}" '
        'alt="Social First Lab · portafolio de marcas">'
        if cover_data else ""
    )
    kpis = "".join(
        f"<article class=\"kpi\"><span>{html.escape(item['label'])}</span><strong>{html.escape(item['value'])}</strong><small>{html.escape(item['context'])} · <a href=\"#slide-{item['slide']}\">diapositiva {item['slide']}</a></small></article>"
        for item in analysis["kpis"] if item.get("label") and item.get("value")
    )
    findings = "".join(
        f"<article class=\"finding\"><span class=\"slide-index\">HALLAZGO {index:02d}</span><h3>{html.escape(item['title'])}</h3><p>{html.escape(item['detail'])}</p><p class=\"evidence\">{html.escape(item['evidence'])} · {', '.join(f'<a href=\"#slide-{n}\">diapositiva {n}</a>' for n in item['slides'])}</p></article>"
        for index, item in enumerate(analysis["findings"], 1)
    )
    actions = "".join(f"<li>{html.escape(action)}</li>" for action in analysis["actions"])
    limitations = "".join(f"<li>{html.escape(note)}</li>" for note in analysis["limitations"])
    slides_html = "".join(_slide_html(slide, deck["assets"]) for slide in deck["slides"])
    metadata = f"{deck['slide_count']} diapositivas · {deck['unique_image_count']} imágenes únicas · {deck['table_count']} tablas · {deck['chart_count']} gráficos"
    return f"""<!doctype html>
<html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="color-scheme" content="light"><title>{html.escape(title)} · Resumen ejecutivo</title>
<style>
 :root{{--ink:#132251;--muted:#596b8c;--line:#dbe4f1;--paper:#f4f7fc;--brand:#2438c8;--soft:#eaf0fb;--green:#087e65}}*{{box-sizing:border-box}}html{{scroll-behavior:smooth}}body{{margin:0;background:var(--paper);color:var(--ink);font:16px/1.55 Inter,ui-sans-serif,system-ui,-apple-system,"Segoe UI",sans-serif}}main{{max-width:1180px;margin:auto;padding:clamp(20px,5vw,64px)}}.hero{{padding:clamp(26px,5vw,60px);background:#fff;border:1px solid var(--line);border-radius:24px;box-shadow:0 18px 48px #17245e0c}}.brand-cover{{display:block;width:100%;height:auto;aspect-ratio:16/9;object-fit:cover;border-radius:14px;margin:0 0 28px}}.eyebrow,.slide-index{{font-size:11px;font-weight:800;letter-spacing:.14em;color:var(--brand);text-transform:uppercase}}h1{{font-size:clamp(32px,6vw,58px);line-height:1.03;letter-spacing:-.045em;margin:12px 0}}.subtitle,.muted{{color:var(--muted)}}.meta{{font-size:13px;color:var(--muted);margin-top:22px}}.summary{{font-size:clamp(18px,2vw,23px);max-width:900px;margin:28px 0 0}}.kpis{{display:grid;grid-template-columns:repeat(auto-fit,minmax(190px,1fr));gap:12px;margin:18px 0 44px}}.kpi{{background:#fff;border:1px solid var(--line);border-radius:16px;padding:18px}}.kpi span,.kpi small{{display:block;color:var(--muted);font-size:12px}}.kpi strong{{display:block;font-size:30px;letter-spacing:-.03em;margin:5px 0}}a{{color:var(--brand)}}.section-head{{margin:52px 0 14px}}.section-head h2{{font-size:clamp(25px,4vw,36px);letter-spacing:-.035em;margin:5px 0}}.findings{{display:grid;grid-template-columns:repeat(auto-fit,minmax(260px,1fr));gap:14px}}.finding{{background:#fff;border:1px solid var(--line);border-radius:16px;padding:22px}}.finding h3{{font-size:19px;margin:8px 0}}.finding p{{margin:8px 0}}.evidence{{font-size:13px;color:var(--muted)}}.actions,.limitations{{background:#fff;border-left:3px solid var(--brand);padding:16px 24px;margin:14px 0;border-radius:0 14px 14px 0}}.slides-nav{{display:flex;gap:8px;overflow:auto;padding:8px 0 16px}}.slides-nav a{{flex:0 0 auto;padding:8px 12px;background:#fff;border:1px solid var(--line);border-radius:99px;text-decoration:none;font-size:13px}}.slide{{background:#fff;border:1px solid var(--line);border-radius:18px;padding:clamp(18px,3vw,30px);margin:16px 0;scroll-margin-top:20px}}.slide h2{{font-size:clamp(20px,3vw,28px);margin:5px 0 20px}}.slide-copy ul{{padding-left:20px}}.slide-copy li{{margin:6px 0;white-space:pre-wrap}}.slide-data{{margin:22px 0;padding:16px;background:#f7f9fd;border-radius:12px}}.slide-data h3,.slide-copy h3,.slide-visuals h3{{font-size:15px;margin:0 0 10px}}.slide-data p{{margin:5px 0;overflow-wrap:anywhere}}.table-scroll{{overflow:auto}}table{{border-collapse:collapse;min-width:100%}}td{{padding:8px;border-bottom:1px solid var(--line);text-align:left}}.image-grid{{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(100%,220px),1fr));gap:14px}}figure{{margin:0;padding:10px;border:1px solid var(--line);border-radius:12px;background:#fff}}figure img{{display:block;max-width:100%;max-height:520px;object-fit:contain;margin:auto}}figcaption{{font-size:11px;color:var(--muted);padding-top:8px}}footer{{padding:24px 0;color:var(--muted);font-size:12px;border-top:1px solid var(--line);margin-top:42px}}@media print{{body{{background:#fff}}main{{max-width:none;padding:0}}.hero,.slide,.finding,.kpi{{box-shadow:none;break-inside:avoid}}.slides-nav{{display:none}}a{{color:inherit;text-decoration:none}}}}
</style></head><body><main>
<header class="hero">{cover_html}<div class="eyebrow">REPORTE DE PRESENTACIÓN</div><h1>{html.escape(title)}</h1><div class="subtitle">{html.escape(deck['filename'])}</div><p class="summary">{html.escape(analysis['executive_summary'])}</p><p class="meta">{html.escape(metadata)} · Las imágenes originales se incluyen como referencia.</p></header>
{f'<section class="kpis" aria-label="Indicadores extraídos">{kpis}</section>' if kpis else ''}
{f'<section><div class="section-head"><div class="eyebrow">LECTURA EJECUTIVA</div><h2>Hallazgos principales</h2></div><div class="findings">{findings}</div></section>' if findings else ''}
{f'<section class="actions"><div class="eyebrow">PRÓXIMOS PASOS</div><ul>{actions}</ul></section>' if actions else ''}
{f'<section class="limitations"><div class="eyebrow">ALCANCE Y LÍMITES</div><ul>{limitations}</ul></section>' if limitations else ''}
<section><div class="section-head"><div class="eyebrow">FUENTE COMPLETA</div><h2>Contenido por diapositiva</h2><p class="subtitle">Cada hallazgo enlaza con las diapositivas usadas como evidencia.</p></div><nav class="slides-nav" aria-label="Navegar diapositivas">{''.join(f'<a href="#slide-{slide["number"]}">Diapositiva {slide["number"]:02d}</a>' for slide in deck['slides'])}</nav>{slides_html}</section>
<footer>Reporte generado desde el contenido de la presentación. Los hallazgos se vinculan con diapositivas fuente. El contenido visual se conserva como imagen original y no se describe automáticamente.</footer>
</main></body></html>"""


def create_report(data: bytes, filename: str) -> str:
    deck = extract_presentation(data, filename)
    return build_report_html(deck)
