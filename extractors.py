import csv, io, json
from pathlib import Path

def extract(data: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True); tables = []
        for ws in wb.worksheets:
            tables.append({"sheet": ws.title, "rows": [[x for x in row] for row in ws.iter_rows(values_only=True)]})
        return "__XLSX_JSON__" + json.dumps(tables, ensure_ascii=False, default=str)
    if ext == ".csv": return "\n".join(" | ".join(row) for row in csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data)); out = []
        for i, slide in enumerate(prs.slides, 1):
            out.append(f"Diapositiva {i}"); out.extend(shape.text for shape in slide.shapes if hasattr(shape, "text"))
        return "\n".join(out)
    raise ValueError("Formato no soportado. Usa PDF, XLSX, CSV o PPTX.")
